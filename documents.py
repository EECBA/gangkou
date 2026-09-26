# -*- coding: utf-8 -*-
"""documents.py —— 她的文档工坊（2026.9.19）：她给他做 Word / PPT。

设计约定：
- 模型（她）只产出一份结构化的文档规格（JSON），渲染全在本模块——
  不让模型写代码也不让它拼 XML，格式可控、坏规格有可读报错。
- 文件落 WORKSPACE_DIR（soulhome/她做的文件/），文件名去非法字符、
  撞名自动加序号；聊天里以文件卡片交付（server.py 挂 files 字段）。
- 内容总量封顶 CONTENT_CAP 字：防她一口气写超长文档把 token 撑爆
  （超了截断并在返回里说明，让她分次做）。
改渲染样式来这里；工具的接线和提示词在 server.py。
"""
import re
import time
from pathlib import Path

from fastapi import APIRouter
from fastapi.responses import FileResponse

from common import WORKSPACE_DIR

router = APIRouter()

CONTENT_CAP = 15000          # 全文档正文字数封顶（超长截断，提示分次）
DOC_DEFAULTS = {"font": "宋体", "size_pt": 12, "first_line_indent_chars": 2,
                "line_spacing": 1.5, "align": "justify"}
PPT_THEMES = {
    # frame=细线框颜色 frame2=内框颜色（None=单线）——纸感版式二期 2026.9.19
    "paper": {"bg": "F6F0DF", "ink": "2C2A26", "sub": "6B655A", "accent": "B08D3E", "label": "米黄纸面",
              "frame": "2C2A26", "frame2": "B08D3E"},
    "dark":  {"bg": "23262E", "ink": "EDE7D8", "sub": "9AA0AC", "accent": "D4B463", "label": "墨夜",
              "frame": "D4B463", "frame2": None},
    "clean": {"bg": "FFFFFF", "ink": "1F2328", "sub": "6E7480", "accent": "8A6D2F", "label": "素白",
              "frame": "9AA0AC", "frame2": None},
}
PPT_DEFAULTS = {"font": "微软雅黑", "title_size_pt": 30, "body_size_pt": 17, "theme": "clean"}   # 默认素白（用户 2026.9.19 三选一拍板）


def _clamp_num(v, lo, hi, dft):
    try:
        v = float(v)
    except (TypeError, ValueError):
        return dft
    return v if lo <= v <= hi else dft


def _sanitize_stem(raw: str, fallback: str) -> str:
    s = re.sub(r'[\\/:*?"<>|\r\n\t]', "", str(raw or "")).strip().strip(".")
    s = s[:60].strip()
    return s or (re.sub(r'[\\/:*?"<>|\r\n\t]', "", fallback)[:20].strip() or "未命名")


def _unique_path(stem: str, ext: str) -> Path:
    cand = WORKSPACE_DIR / f"{stem}{ext}"
    n = 2
    while cand.exists():
        cand = WORKSPACE_DIR / f"{stem}({n}){ext}"
        n += 1
    return cand


def _block_text_total(spec: dict) -> int:
    total = 0
    for b in spec.get("blocks") or []:
        if isinstance(b, dict):
            total += len(str(b.get("text") or ""))
            total += sum(len(str(x)) for x in (b.get("items") or []) if isinstance(x, str))
    return total


def _slide_text_total(spec: dict) -> int:
    total = 0
    for s in spec.get("slides") or []:
        if isinstance(s, dict):
            total += len(str(s.get("title") or "")) + len(str(s.get("subtitle") or "")) + len(str(s.get("body") or ""))
            total += sum(len(str(x)) for x in (s.get("items") or []) if isinstance(x, str))
    return total


# ---------------- Word ----------------

def render_docx(spec: dict, out_dir: Path = None) -> dict:
    """把规格渲染成 .docx。返回 {ok, name, size} 或 {ok:False, error}。"""
    blocks = spec.get("blocks")
    if not isinstance(blocks, list) or not blocks:
        return {"ok": False, "error": "blocks 是空的——把要写的完整内容放进 blocks（heading/paragraph/list/quote）"}
    st = dict(DOC_DEFAULTS)
    try:
        st.update({k: v for k, v in (spec.get("styles") or {}).items() if v is not None})
    except Exception:
        pass
    font = str(st["font"])[:30] or "宋体"
    size = _clamp_num(st["size_pt"], 8, 36, 12)
    indent = _clamp_num(st["first_line_indent_chars"], 0, 8, 2)
    spacing = _clamp_num(st["line_spacing"], 1.0, 3.0, 1.5)
    align = {"left": 0, "center": 1, "right": 2, "justify": 3}.get(str(st["align"]), 3)

    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document()

    def set_font(run, size_pt, bold=False, color=None):
        run.font.name = font
        run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
        run.font.size = Pt(size_pt)
        run.font.bold = bold
        if color:
            run.font.color.rgb = RGBColor.from_string(color)

    normal = doc.styles["Normal"]
    normal.font.name = font
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    normal.font.size = Pt(size)
    npf = normal.paragraph_format
    npf.first_line_indent = Pt(size * indent)
    npf.line_spacing = spacing
    npf.alignment = WD_ALIGN_PARAGRAPH(align)

    title = str(spec.get("title") or "").strip()
    if title:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(size * 0.8)
        set_font(p.add_run(title[:100]), size * 1.6, bold=True)

    used, overflow = 0, False
    for b in blocks:
        if not isinstance(b, dict):
            continue
        kind = str(b.get("type") or "paragraph")
        text = str(b.get("text") or "").strip()
        items = [str(x) for x in (b.get("items") or []) if str(x).strip()]
        if kind != "list" and not text:
            continue
        chunk = len(text) + sum(len(x) for x in items)
        if used + chunk > CONTENT_CAP:
            overflow = True
            break
        used += chunk
        if kind == "heading":
            level = int(_clamp_num(b.get("level"), 1, 3, 1))
            sc = {1: 1.4, 2: 1.2, 3: 1.05}[level]
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.first_line_indent = Pt(0)
            pf.space_before = Pt(size * 0.6)
            pf.space_after = Pt(size * 0.3)
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
            set_font(p.add_run(text[:200]), size * sc, bold=True)
        elif kind == "list":
            ordered = bool(b.get("ordered"))
            for i, item in enumerate(items, 1):
                p = doc.add_paragraph()
                pf = p.paragraph_format
                pf.first_line_indent = Pt(0)
                pf.left_indent = Pt(size * 1.0)
                set_font(p.add_run(f"{i}. {item}" if ordered else f"• {item}"), size)
        elif kind == "quote":
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.first_line_indent = Pt(0)
            pf.left_indent = Pt(size * 2)
            set_font(p.add_run(text[:800]), size, color="6B655A")
        else:
            p = doc.add_paragraph()
            set_font(p.add_run(text), size)

    stem = _sanitize_stem(spec.get("filename"), title or "文档")
    out = _unique_path(stem, ".docx")
    doc.save(str(out))
    note = "；内容到字数上限，后面的没写进去，可以让她分次续做" if overflow else ""
    return {"ok": True, "name": out.name, "size": out.stat().st_size, "note": note}


# ---------------- PPT ----------------

def render_pptx(spec: dict, out_dir: Path = None) -> dict:
    """把规格渲染成 .pptx（16:9，无模板，手排版式——一期朴素但干净）。"""
    slides = spec.get("slides")
    if not isinstance(slides, list) or not slides:
        return {"ok": False, "error": "slides 是空的——每一页一个 {type, title, ...}，封面用 cover、收尾用 end"}
    st = dict(PPT_DEFAULTS)
    try:
        st.update({k: v for k, v in (spec.get("styles") or {}).items() if v is not None})
    except Exception:
        pass
    font = str(st["font"])[:30] or "微软雅黑"
    tsize = _clamp_num(st["title_size_pt"], 16, 60, 30)
    bsize = _clamp_num(st["body_size_pt"], 10, 40, 17)
    theme = PPT_THEMES.get(str(st["theme"]), PPT_THEMES["paper"])

    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    SW, SH = 13.333, 7.5
    FR = 0.42            # 外框内缩
    FR2 = 0.56           # 内框内缩（双线主题）

    def new_slide():
        s = prs.slides.add_slide(blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string(theme["bg"])

        def frame(inset, color, w_pt):
            sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(inset), Inches(inset),
                                    Inches(SW - inset * 2), Inches(SH - inset * 2))
            sh.fill.background()
            sh.line.color.rgb = RGBColor.from_string(color)
            sh.line.width = Pt(w_pt)
            sh.shadow.inherit = False
            return sh

        frame(FR, theme["frame"], 1.0)
        if theme.get("frame2"):
            frame(FR2, theme["frame2"], 0.75)
        return s

    def box(slide, l, t, w, h, anchor=MSO_ANCHOR.TOP):
        tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        return tf

    def para(tf, text, size_pt, color, bold=False, align=PP_ALIGN.LEFT, first=False, space_after=6):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        r = p.add_run()
        r.text = str(text)
        r.font.name = font
        r.font.size = Pt(size_pt)
        r.font.bold = bold
        r.font.color.rgb = RGBColor.from_string(color)
        return p

    def gold_bar(slide, l, t, w=1.6, h=3.2):
        sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Pt(h))
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(theme["accent"])
        sh.line.fill.background()
        sh.shadow.inherit = False

    def footer(slide, no, total, deck_title, left=True):
        if left and deck_title:
            tf = box(slide, FR + 0.35, SH - FR - 0.5, 8.0, 0.4)
            para(tf, deck_title[:40], 9.5, theme["sub"], first=True)
        tf = box(slide, SW - FR - 1.45, SH - FR - 0.5, 1.1, 0.4)
        para(tf, f"{no} / {total}", 9.5, theme["sub"], align=PP_ALIGN.RIGHT, first=True)

    total = len(slides)
    used, overflow = 0, False
    kept = []
    for s in slides:
        if isinstance(s, dict):
            chunk = (len(str(s.get("title") or "")) + len(str(s.get("subtitle") or ""))
                     + len(str(s.get("body") or "")) + sum(len(str(x)) for x in (s.get("items") or [])))
            if used + chunk > CONTENT_CAP:
                overflow = True
                break
            used += chunk
            kept.append(s)

    deck_title = str(spec.get("title") or "").strip()
    LX, CW = FR + 0.5, SW - 2 * (FR + 0.5)   # 正文左基线与内容宽
    for no, s in enumerate(kept, 1):
        kind = str(s.get("type") or "bullets")
        slide = new_slide()
        if kind == "cover":
            dia = box(slide, LX, 1.95, CW, 0.5)
            para(dia, "◆", 15, theme["accent"], align=PP_ALIGN.CENTER, first=True)
            tf = box(slide, LX, 2.5, CW, 1.6)
            para(tf, str(s.get("title") or deck_title or ""), tsize + 10, theme["ink"], bold=True,
                 align=PP_ALIGN.CENTER, first=True, space_after=10)
            sub = str(s.get("subtitle") or "").strip()
            if sub:
                para(tf, sub, bsize, theme["sub"], align=PP_ALIGN.CENTER)
            d = box(slide, SW - FR - 2.3, SH - FR - 0.55, 1.9, 0.4)
            para(d, time.strftime("%Y年%m月%d日"), 11, theme["sub"], align=PP_ALIGN.RIGHT, first=True)
        elif kind == "section":
            num = box(slide, LX, 2.15, CW, 0.9)
            para(num, f"{no:02d}", tsize * 0.9, theme["accent"], bold=True,
                 align=PP_ALIGN.CENTER, first=True)
            tf = box(slide, LX, 3.15, CW, 1.4)
            para(tf, str(s.get("title") or ""), tsize + 2, theme["ink"], bold=True,
                 align=PP_ALIGN.CENTER, first=True)
            footer(slide, no, total, deck_title)
        elif kind == "end":
            tf = box(slide, LX, 2.9, CW, 1.3)
            para(tf, str(s.get("title") or "谢谢观看"), tsize + 4, theme["ink"], bold=True,
                 align=PP_ALIGN.CENTER, first=True)
            dia = box(slide, LX, 4.35, CW, 0.5)
            para(dia, "◆", 13, theme["accent"], align=PP_ALIGN.CENTER, first=True)
            footer(slide, no, total, deck_title, left=False)
        else:  # bullets / text
            title = str(s.get("title") or "").strip()
            top = 0.82
            if title:
                tf = box(slide, LX, top, CW, 0.85)
                para(tf, title, tsize, theme["ink"], bold=True, first=True)
                gold_bar(slide, LX, top + 0.92, w=1.5)
            items = [str(x) for x in (s.get("items") or []) if str(x).strip()]
            body = str(s.get("body") or "").strip()
            tf = box(slide, LX, 2.12, CW, 4.3)
            first = True
            if items:
                for it in items:
                    para(tf, f"·  {it}", bsize, theme["ink"], first=first, space_after=9)
                    first = False
            elif body:
                for seg in body.split("\n"):
                    if seg.strip():
                        para(tf, seg.strip(), bsize, theme["ink"], first=first, space_after=9)
                        first = False
            footer(slide, no, total, deck_title)

    stem = _sanitize_stem(spec.get("filename"), str(spec.get("title") or "演示"))
    out = _unique_path(stem, ".pptx")
    prs.save(str(out))
    note = "；页数到字数上限，后面的页没做进去，可以让她分次续做" if overflow else ""
    return {"ok": True, "name": out.name, "size": out.stat().st_size, "note": note}


# ---------------- 文本抽取（他发来的文件，她用 read_file 读） ----------------

def extract_text(path: Path) -> str:
    """把一个文件抽成纯文本给她读。txt/md/csv 直读；docx 走 python-docx（含表格）；
    pdf 走 pypdf。抽不出来就抛异常，调用方转成可读报错。"""
    sfx = path.suffix.lower()
    if sfx == ".docx":
        from docx import Document
        d = Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    if sfx == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        return "\n".join((pg.extract_text() or "").strip() for pg in reader.pages).strip()
    return path.read_text(encoding="utf-8", errors="ignore")


# ---------------- 下载路由 ----------------

@router.get("/workspace/{name}")
def download_workspace_file(name: str):
    """聊天里的文件卡片指向这里；只允许从工作区里取，防路径穿越。"""
    from fastapi.responses import JSONResponse
    safe = re.sub(r'[\\/:*?"<>|\r\n\t]', "", name).strip()
    p = (WORKSPACE_DIR / safe).resolve()
    if not str(p).startswith(str(WORKSPACE_DIR.resolve())) or not p.is_file():
        return JSONResponse({"error": "没有这个文件"}, status_code=404)
    return FileResponse(p, filename=p.name)
